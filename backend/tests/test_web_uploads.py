from __future__ import annotations

import io
import json
import os
import zipfile
from datetime import datetime, timedelta, timezone

import pytest

from app.web_api.document_extraction import DocumentValidationError
from app.web_api.attachment_context import select_attachment_context
from app.web_api.upload_store import (
    EphemeralUpload, ExtractedChunk, RedisEphemeralUploadStore, get_upload_store,
    reset_upload_store_for_tests, upload_ttl_seconds, utc_iso,
)
from tests.conftest import auth_headers, create_test_user


MIME = {
    "txt": "text/plain",
    "md": "text/markdown",
    "csv": "text/csv",
    "json": "application/json",
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def _upload(client, filename: str, content: bytes, media_type: str, *, uid: str = "upload-user"):
    return client.post(
        "/api/web/uploads",
        headers=auth_headers(uid, f"{uid}@example.com"),
        files={"file": (filename, content, media_type)},
    )


def _supported_files() -> dict[str, bytes]:
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from reportlab.pdfgen import canvas

    pdf = io.BytesIO()
    pdf_canvas = canvas.Canvas(pdf)
    pdf_canvas.drawString(72, 720, "PDF source text")
    pdf_canvas.save()

    docx = io.BytesIO()
    document = Document()
    document.add_paragraph("DOCX paragraph text")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "DOCX table text"
    document.save(docx)

    xlsx = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Item", "Amount"])
    sheet.append(["Hosting", 42])
    workbook.save(xlsx)

    pptx = io.BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "PPTX slide text"
    presentation.save(pptx)
    return {
        "txt": b"First line\nSecond line",
        "md": b"# Heading\nMarkdown body",
        "csv": b"name,total\nHosting,42\n",
        "json": b'{"project":"Swico","total":42}',
        "pdf": pdf.getvalue(),
        "docx": docx.getvalue(),
        "xlsx": xlsx.getvalue(),
        "pptx": pptx.getvalue(),
    }


def _office_zip(extension: str, extra: dict[str, bytes]) -> bytes:
    expected = {
        "docx": "word/document.xml",
        "xlsx": "xl/workbook.xml",
        "pptx": "ppt/presentation.xml",
    }[extension]
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr(expected, b"<document/>")
        for name, content in extra.items():
            archive.writestr(name, content)
    return output.getvalue()


def test_upload_requires_authentication(client):
    response = client.post("/api/web/uploads", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert response.status_code == 401


def test_unsupported_mismatch_empty_and_per_file_limit(client, monkeypatch):
    create_test_user("upload-user", "upload-user@example.com")
    assert _upload(client, "script.py", b"print('x')", "text/x-python").json()["error"]["code"] == "unsupported_file_type"
    mismatch = _upload(client, "notes.txt", b"hello", "application/pdf")
    assert mismatch.status_code == 422 and mismatch.json()["error"]["code"] == "media_type_mismatch"
    empty = _upload(client, "notes.txt", b"", "text/plain")
    assert empty.status_code == 400 and empty.json()["error"]["code"] == "empty_file"
    monkeypatch.setenv("WEB_UPLOAD_MAX_FILE_BYTES", "4")
    too_large = _upload(client, "notes.txt", b"12345", "text/plain")
    assert too_large.status_code == 413 and too_large.json()["error"]["code"] == "file_too_large"


def test_octet_stream_requires_a_valid_signature_and_legacy_doc_is_honest(client, monkeypatch):
    create_test_user("upload-user", "upload-user@example.com")
    monkeypatch.setenv("WEB_LEGACY_DOC_CONVERSION_ENABLED", "false")
    valid = _upload(client, "report.pdf", _supported_files()["pdf"], "application/octet-stream")
    assert valid.status_code == 201
    invalid = _upload(client, "report.pdf", b"not a pdf", "application/octet-stream")
    assert invalid.status_code == 422
    assert invalid.json()["error"]["code"] == "invalid_pdf"
    legacy = _upload(client, "report.doc", b"\xd0\xcf\x11\xe0", "application/msword")
    assert legacy.status_code == 422
    assert legacy.json()["error"]["code"] == "legacy_doc_unsupported"
    assert "re-save" in legacy.json()["error"]["message"].lower()
    assert ".docx" in legacy.json()["error"]["message"].lower()


def test_encrypted_pdf_is_rejected_explicitly(client):
    from pypdf import PdfWriter

    create_test_user("upload-user", "upload-user@example.com")
    value = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("private-password")
    writer.write(value)
    response = _upload(client, "encrypted.pdf", value.getvalue(), MIME["pdf"])
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "encrypted_document"


def test_filename_is_sanitized_and_macro_extension_rejected(client):
    create_test_user("upload-user", "upload-user@example.com")
    response = _upload(client, "../../private/notes.txt", b"safe", "text/plain")
    assert response.status_code == 201
    assert response.json()["name"] == "notes.txt"
    macro = _upload(client, "unsafe.docm", b"zip", "application/vnd.ms-word.document.macroEnabled.12")
    assert macro.status_code == 422 and macro.json()["error"]["code"] == "unsupported_file_type"


def test_malicious_office_paths_and_zip_bombs_are_rejected(client):
    create_test_user("upload-user", "upload-user@example.com")
    traversal = _office_zip("docx", {"../payload.exe": b"bad"})
    response = _upload(client, "bad.docx", traversal, MIME["docx"])
    assert response.status_code == 422 and response.json()["error"]["code"] == "unsafe_office_container"
    bomb = _office_zip("docx", {"word/big.xml": b"A" * 2_000_000})
    response = _upload(client, "bomb.docx", bomb, MIME["docx"])
    assert response.status_code == 422 and response.json()["error"]["code"] == "zip_bomb_detected"


@pytest.mark.parametrize("extension", ["txt", "md", "csv", "json", "pdf", "docx", "xlsx", "pptx"])
def test_successful_extraction_for_every_supported_format_preserves_sources(client, extension):
    create_test_user("upload-user", "upload-user@example.com")
    response = _upload(client, f"sample.{extension}", _supported_files()[extension], MIME[extension])
    assert response.status_code == 201, response.text
    payload = response.json()
    upload = get_upload_store().get(payload["id"])
    assert upload is not None and upload.chunks
    assert upload.source_locators
    expected_source = {
        "txt": "lines ", "md": "lines ", "csv": "row ", "json": "lines ",
        "pdf": "page ", "docx": "paragraph ", "xlsx": "Budget row ", "pptx": "slide ",
    }[extension]
    assert any(source.startswith(expected_source) for source in upload.source_locators)


def test_docx_extracts_headers_and_footers(client):
    from docx import Document
    create_test_user("upload-user", "upload-user@example.com")
    value = io.BytesIO()
    document = Document()
    document.add_paragraph("Body")
    document.sections[0].header.paragraphs[0].text = "Header text"
    document.sections[0].footer.paragraphs[0].text = "Footer text"
    document.save(value)
    response = _upload(client, "header.docx", value.getvalue(), MIME["docx"])
    assert response.status_code == 201
    upload = get_upload_store().get(response.json()["id"])
    assert upload is not None
    assert any("header" in chunk.source and "Header text" in chunk.text for chunk in upload.chunks)
    assert any("footer" in chunk.source and "Footer text" in chunk.text for chunk in upload.chunks)


def test_zero_text_pdf_is_ready_with_explicit_no_ocr_warning(client, monkeypatch):
    from reportlab.pdfgen import canvas
    monkeypatch.setenv("WEB_DOCUMENT_OCR_ENABLED", "false")
    create_test_user("upload-user", "upload-user@example.com")
    value = io.BytesIO()
    document = canvas.Canvas(value)
    document.showPage(); document.save()
    response = _upload(client, "scan.pdf", value.getvalue(), MIME["pdf"])
    assert response.status_code == 201
    payload = response.json()
    assert "pdf_no_extractable_text" in payload["warning_codes"]
    assert "This PDF looks scanned — no selectable text was found." in payload["warnings"]
    upload = get_upload_store().get(payload["id"])
    assert upload is not None and upload.chunks == []


def test_virtual_text_upload_accepts_50000_chars_and_chunks(client, monkeypatch):
    monkeypatch.setenv("WEB_LONG_INPUT_ENABLED", "true")
    create_test_user("upload-user", "upload-user@example.com")
    text = ("architecture roadmap database api testing deployment " * 1000)[:50_000]
    response = client.post(
        "/api/web/uploads/text",
        headers=auth_headers("upload-user", "upload-user@example.com"),
        json={
            "upload_id": "90000000-0000-4000-8000-000000000050",
            "text": text, "operation": "analyze",
        },
    )
    assert response.status_code == 201
    upload = get_upload_store().get(response.json()["id"])
    assert upload is not None and len(upload.chunks) > 20
    assert upload.virtual_text_operation == "analyze"
    assert max(len(chunk.text) for chunk in upload.chunks) <= 2000
    provider_excerpt = select_attachment_context([upload], "Which deployment architecture is described?")
    assert len(provider_excerpt) <= 8000
    assert len(provider_excerpt) < len(text)
    full_operation = client.post(
        "/api/web/chat/stream",
        headers=auth_headers("upload-user", "upload-user@example.com"),
        json={
            "request_id": "90000000-0000-4000-8000-000000000051",
            "message": "Analyze the attached pasted text. Preserve its meaning.",
            "attachment_ids": [upload.id],
        },
    )
    assert full_operation.status_code == 422
    assert full_operation.json()["error"]["code"] == "full_document_confirmation_required"


def test_upload_ownership_isolation_explicit_delete_and_expiry(client):
    create_test_user("upload-user", "upload-user@example.com")
    create_test_user("other-user", "other-user@example.com")
    uploaded = _upload(client, "notes.txt", b"private facts", "text/plain").json()
    forbidden = client.delete(f"/api/web/uploads/{uploaded['id']}", headers=auth_headers("other-user", "other-user@example.com"))
    assert forbidden.status_code == 404
    assert get_upload_store().get(uploaded["id"]) is not None
    deleted = client.delete(f"/api/web/uploads/{uploaded['id']}", headers=auth_headers("upload-user", "upload-user@example.com"))
    assert deleted.status_code == 204 and get_upload_store().get(uploaded["id"]) is None
    assert client.delete(f"/api/web/uploads/{uploaded['id']}", headers=auth_headers("upload-user", "upload-user@example.com")).status_code == 204


def test_total_attachment_limit_and_expired_attachment_return_before_billing(client, monkeypatch):
    create_test_user("upload-user", "upload-user@example.com")
    first = _upload(client, "one.txt", b"first", "text/plain").json()
    second = _upload(client, "two.txt", b"second", "text/plain").json()
    monkeypatch.setenv("WEB_UPLOAD_MAX_TOTAL_BYTES", "8")
    total = client.post("/api/web/chat/stream", headers=auth_headers("upload-user", "upload-user@example.com"), json={
        "request_id": "90000000-0000-4000-8000-000000000001", "message": "summarize",
        "attachment_ids": [first["id"], second["id"]],
    })
    assert total.status_code == 413 and total.json()["error"]["code"] == "attachment_total_too_large"
    missing = client.post("/api/web/chat/stream", headers=auth_headers("upload-user", "upload-user@example.com"), json={
        "request_id": "90000000-0000-4000-8000-000000000002", "message": "summarize",
        "attachment_ids": ["00000000-0000-4000-8000-000000000099"],
    })
    assert missing.status_code == 410 and missing.json()["error"]["code"] == "attachment_expired"


def test_one_user_cannot_use_another_users_attachment(client):
    create_test_user("upload-user", "upload-user@example.com")
    create_test_user("other-user", "other-user@example.com")
    uploaded = _upload(client, "notes.txt", b"private facts", "text/plain").json()
    response = client.post("/api/web/chat/stream", headers=auth_headers("other-user", "other-user@example.com"), json={
        "request_id": "90000000-0000-4000-8000-000000000003", "message": "reveal it",
        "attachment_ids": [uploaded["id"]],
    })
    assert response.status_code == 404 and response.json()["error"]["code"] == "attachment_not_found"


def test_redis_setex_uses_600_and_reads_do_not_renew(monkeypatch):
    calls: list[tuple] = []
    values: dict[str, str] = {}

    class FakeRedis:
        def setex(self, key, ttl, value): calls.append(("setex", key, ttl)); values[key] = value
        def get(self, key): calls.append(("get", key)); return values.get(key)
        def exists(self, key): calls.append(("exists", key)); return int(key in values)
        def delete(self, key): values.pop(key, None); return 1
        def ping(self): return True

    monkeypatch.setattr("redis.Redis.from_url", lambda *args, **kwargs: FakeRedis())
    store = RedisEphemeralUploadStore("redis://private", ttl_seconds=600)
    upload = EphemeralUpload(
        id="00000000-0000-4000-8000-000000000010", owner_user_id=1, name="notes.txt",
        extension=".txt", media_type="text/plain", size_bytes=5, created_at=utc_iso(),
        expires_at=utc_iso(datetime.now(timezone.utc) + timedelta(minutes=10)),
        chunks=[ExtractedChunk("hello", "line 1")], source_locators=["line 1"], warnings=[],
    )
    store.put(upload)
    assert store.is_owned(upload.id, 1) is True
    assert store.is_owned(upload.id, 2) is False
    store.get(upload.id)
    assert calls[0][0] == "setex" and calls[0][2] == 600
    assert calls[1][0] == "setex" and calls[1][2] == 600
    assert [call[0] for call in calls] == [
        "setex", "setex", "exists", "exists", "get",
    ]


def test_upload_ttl_defaults_to_one_hour_and_is_capped_at_24_hours(monkeypatch):
    monkeypatch.delenv("WEB_UPLOAD_TTL_SECONDS", raising=False)
    assert upload_ttl_seconds() == 3600
    monkeypatch.setenv("WEB_UPLOAD_TTL_SECONDS", "999999")
    assert upload_ttl_seconds() == 86400


@pytest.mark.parametrize("parser_fails", [False, True])
def test_raw_temporary_file_removed_after_parser_success_or_exception(client, monkeypatch, parser_fails):
    create_test_user("upload-user", "upload-user@example.com")
    import app.web_api.router as router

    original_temp = router.tempfile.NamedTemporaryFile
    paths: list[str] = []

    def tracked_temp(*args, **kwargs):
        handle = original_temp(*args, **kwargs)
        paths.append(handle.name)
        return handle

    monkeypatch.setattr(router.tempfile, "NamedTemporaryFile", tracked_temp)
    if parser_fails:
        monkeypatch.setattr(router, "extract_document", lambda *args: (_ for _ in ()).throw(
            DocumentValidationError("document_parse_failed", "The document could not be parsed safely.")
        ))
    response = _upload(client, "notes.txt", b"temporary", "text/plain")
    assert response.status_code == (422 if parser_fails else 201)
    assert paths and all(not os.path.exists(path) for path in paths)


def test_production_fails_closed_without_upload_cache_url(client, monkeypatch):
    create_test_user("upload-user", "upload-user@example.com")
    monkeypatch.setenv("APP_ENV", "production")
    monkeypatch.delenv("WEB_UPLOAD_CACHE_URL", raising=False)
    monkeypatch.setattr("app.auth.validate_auth_configuration", lambda: None)
    reset_upload_store_for_tests()
    response = _upload(client, "notes.txt", b"temporary", "text/plain")
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "attachment_cache_unavailable"
    bootstrap = client.get("/api/web/bootstrap", headers=auth_headers("upload-user", "upload-user@example.com"))
    assert bootstrap.json()["features"]["web_attachments"] is False
    assert bootstrap.json()["uploads"]["available"] is False
