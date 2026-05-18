from __future__ import annotations

from datetime import timedelta

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from sqlmodel import select

from app.ai.orchestrator import run_text_turn
from app.ai.router import AIProviderRouter
from app.ai.tools import classify_folder_category
from app.ai.types import AIRequest
from app.database import SessionLocal
from app.models import DocumentArtifact, Item
from app.time_utils import utc_now
from conftest import auth_headers, create_test_user


class _ExplodingProvider:
    def complete(self, *_args, **_kwargs):
        raise AssertionError("voice tool workflow must not call an LLM provider")


def _request(message: str, user_id: int, *, channel: str = "voice") -> AIRequest:
    return AIRequest(
        user_id=user_id,
        message=message,
        reply_language="ta",
        channel=channel,
        request_id="voice-doc-test",
        metadata={},
    )


def _redirect_docs(monkeypatch, tmp_path):
    import app.main as main_module

    monkeypatch.setattr(main_module, "DOCS_BASE_DIR", tmp_path.resolve())
    monkeypatch.setattr(main_module, "PDF_BASE_DIR", tmp_path.resolve() / "pdf")
    monkeypatch.setattr(main_module, "DOCX_BASE_DIR", tmp_path.resolve() / "docx")
    monkeypatch.setattr(main_module, "EXCEL_BASE_DIR", tmp_path.resolve() / "xlsx")
    monkeypatch.setattr(main_module, "PPT_BASE_DIR", tmp_path.resolve() / "pptx")


def test_folder_classification_is_local_and_deterministic():
    assert classify_folder_category("client lead follow up notes save பண்ணு") == "Business"
    assert classify_folder_category("office meeting points pdf ஆக்கி save பண்ணு") == "Work"
    assert classify_folder_category("அம்மா medicine reminder நாளைக்கு காலை") == "Home"


def test_tamil_voice_intents_route_to_backend_tools():
    cases = [
        ("இந்த meeting points எல்லாம் PDF ஆக்கி Work Folder ல வை", "document"),
        ("நேத்து சொன்ன business notes open பண்ணு", "file_retrieval"),
        ("அம்மா medicine நாளைக்கு காலை remind பண்ணு", "reminder"),
    ]

    for message, expected_intent in cases:
        route = AIProviderRouter().select_route(_request(message, user_id=1))

        assert route.provider == "backend_tool"
        assert route.intent == expected_intent


def test_tamil_voice_document_command_creates_all_formats_from_spoken_source(client, monkeypatch, tmp_path):
    _redirect_docs(monkeypatch, tmp_path)
    create_test_user()
    message = (
        "office meeting points launch checklist budget review "
        "PDF Word Excel PPT ஆக்கி Work folder ல வை"
    )

    response = client.post(
        "/api/chat",
        headers=auth_headers("test-uid", "test@example.com"),
        json={"message": message, "reply_language": "ta"},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["item"]["intent"] == "document"
    assert payload["item"]["category"] == "Work"
    artifacts = payload["meta"]["artifacts"]
    assert {artifact["format"] for artifact in artifacts} == {"pdf", "docx", "xlsx", "pptx"}
    assert all(artifact["download_id"] and artifact["download_url"] for artifact in artifacts)

    expected_body = "office meeting points launch checklist budget review"
    with SessionLocal() as session:
        rows = list(session.exec(select(DocumentArtifact)).all())
    assert len(rows) == 4
    for artifact in rows:
        assert artifact.category == "Work"
        assert artifact.source_text == expected_body
        assert f"{artifact.format}/Work/{artifact.created_at.date().isoformat()}/" in artifact.relative_path
        assert (tmp_path / artifact.relative_path).is_file()

    docx_path = tmp_path / next(row.relative_path for row in rows if row.format == "docx")
    docx_text = "\n".join(paragraph.text for paragraph in Document(str(docx_path)).paragraphs)
    assert expected_body in docx_text
    assert "Created PDF" not in docx_text

    xlsx_path = tmp_path / next(row.relative_path for row in rows if row.format == "xlsx")
    workbook = load_workbook(str(xlsx_path))
    worksheet = workbook.active
    xlsx_text = "\n".join(str(cell.value or "") for row in worksheet.iter_rows() for cell in row)
    assert expected_body in xlsx_text
    assert "Created PDF" not in xlsx_text

    pptx_path = tmp_path / next(row.relative_path for row in rows if row.format == "pptx")
    presentation = Presentation(str(pptx_path))
    pptx_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    assert expected_body in pptx_text
    assert "Created PDF" not in pptx_text


def test_voice_file_retrieval_returns_signed_openable_metadata(client, monkeypatch, tmp_path):
    _redirect_docs(monkeypatch, tmp_path)
    user = create_test_user()
    yesterday = utc_now() - timedelta(days=1)
    relative_path = f"pdf/Business/{yesterday.date().isoformat()}/business_notes.pdf"
    file_path = tmp_path / relative_path
    file_path.parent.mkdir(parents=True, exist_ok=True)
    file_path.write_bytes(b"%PDF-1.4\nbusiness notes\n")
    with SessionLocal() as session:
        item = Item(
            intent="document",
            category="Business",
            raw_text="business notes follow up",
            title="business notes",
            details="Business notes",
            source="voice",
            user_id=user.id,
            created_at=yesterday,
            updated_at=yesterday,
        )
        session.add(item)
        session.commit()
        session.refresh(item)
        session.add(
            DocumentArtifact(
                user_id=int(user.id),
                item_id=int(item.id),
                title="business notes",
                format="pdf",
                category="Business",
                relative_path=relative_path,
                source_text="business notes follow up",
                created_at=yesterday,
            )
        )
        session.commit()

    response = client.post(
        "/api/chat",
        headers=auth_headers("test-uid", "test@example.com"),
        json={"message": "நேத்து சொன்ன business notes open பண்ணு", "reply_language": "ta"},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["item"]["intent"] == "assistant"
    files = payload["meta"]["files"]
    assert files
    assert files[0]["id"]
    assert files[0]["title"] == "business notes"
    assert files[0]["format"] == "pdf"
    assert files[0]["category"] == "Business"
    assert files[0]["relative_path"] == relative_path
    assert files[0]["download_url"]
    assert files[0]["download_id"]
    assert payload["pipeline"]["direct_answer_source"] == "backend_tool"

    download = client.get(files[0]["download_url"], headers=auth_headers("test-uid", "test@example.com"))
    assert download.status_code == 200


def test_voice_tool_routes_do_not_call_providers_and_have_zero_cost():
    user = create_test_user()
    for message in [
        "client follow up note save பண்ணு",
        "office meeting task add பண்ணு",
        "அம்மா medicine நாளைக்கு காலை remind பண்ணு",
        "இந்த meeting points எல்லாம் PDF ஆக்கி Work Folder ல வை",
        "நேத்து சொன்ன business notes open பண்ணு",
    ]:
        with SessionLocal() as session:
            response = run_text_turn(
                session,
                _request(message, int(user.id)),
                existing_context={"openai_provider": _ExplodingProvider(), "sarvam_provider": _ExplodingProvider()},
            )

        assert response.provider == "backend_tool"
        assert response.estimated_cost_amount == 0
