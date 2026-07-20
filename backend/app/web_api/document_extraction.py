from __future__ import annotations

import csv
import io
import json
import os
import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Callable, Iterable

from .upload_store import ExtractedChunk
from .legacy_doc_conversion import configured_legacy_doc_converter, unavailable_message


SUPPORTED_EXTENSIONS = (".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".xlsx", ".pptx")
MEDIA_TYPES: dict[str, set[str]] = {
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain", "text/x-markdown"},
    ".csv": {"text/csv", "application/csv", "text/plain"},
    ".json": {"application/json", "text/json"},
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    ".xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
    ".pptx": {"application/vnd.openxmlformats-officedocument.presentationml.presentation"},
}
OFFICE_EXPECTED_ENTRY = {
    ".docx": "word/document.xml",
    ".xlsx": "xl/workbook.xml",
    ".pptx": "ppt/presentation.xml",
}
FORBIDDEN_UPLOAD_EXTENSIONS = {
    ".docm", ".xlsm", ".pptm", ".zip", ".rar", ".7z", ".tar", ".gz",
    ".exe", ".dll", ".com", ".bat", ".cmd", ".ps1", ".sh", ".js", ".py",
}
FORBIDDEN_ZIP_ENTRY_EXTENSIONS = {
    ".exe", ".dll", ".com", ".bat", ".cmd", ".ps1", ".sh", ".js", ".jar", ".msi",
}
MAX_ZIP_ENTRIES = 2_000
MAX_ZIP_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ZIP_COMPRESSION_RATIO = 100
CHUNK_CHARS = 2_000
CHUNK_OVERLAP_CHARS = 160


class DocumentValidationError(ValueError):
    def __init__(self, code: str, message: str, *, status_code: int = 422) -> None:
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code


@dataclass(frozen=True)
class ExtractionResult:
    chunks: list[ExtractedChunk]
    source_locators: list[str]
    warnings: list[str]
    extracted_chars: int
    warning_codes: list[str]
    page_character_counts: list[int]


def env_int(name: str, default: int, *, minimum: int = 1, maximum: int | None = None) -> int:
    try:
        value = int(os.getenv(name, str(default)))
    except ValueError:
        value = default
    value = max(minimum, value)
    return min(value, maximum) if maximum is not None else value


def max_file_bytes() -> int:
    return env_int("WEB_UPLOAD_MAX_FILE_BYTES", 10 * 1024 * 1024, maximum=10 * 1024 * 1024)


def max_extracted_chars() -> int:
    return env_int("WEB_UPLOAD_MAX_EXTRACTED_CHARS", 100_000, maximum=500_000)


def sanitize_filename(value: str | None) -> str:
    raw = str(value or "document").replace("\\", "/")
    name = raw.rsplit("/", 1)[-1]
    name = "".join(char for char in name if char >= " " and char != "\x7f")
    name = re.sub(r"\s+", " ", name).strip().strip(".")
    if not name:
        name = "document"
    stem = re.sub(r"[^\w .()\[\]-]+", "_", Path(name).stem, flags=re.UNICODE).strip(" ._") or "document"
    suffix = Path(name).suffix.lower()
    return f"{stem[:150]}{suffix[:12]}"


def validate_extension_and_mime(filename: str, content_type: str | None) -> tuple[str, str]:
    extension = Path(filename).suffix.lower()
    if extension == ".doc":
        # The API process intentionally has no native .doc parser. A future
        # isolated worker must implement the converter protocol before this
        # branch may accept the upload.
        configured_legacy_doc_converter()
        raise DocumentValidationError(
            "legacy_doc_conversion_required",
            unavailable_message(),
        )
    if extension in FORBIDDEN_UPLOAD_EXTENSIONS or extension not in SUPPORTED_EXTENSIONS:
        raise DocumentValidationError(
            "unsupported_file_type",
            "This file type is not supported. Upload TXT, MD, CSV, JSON, PDF, DOCX, XLSX, or PPTX.",
        )
    media_type = str(content_type or "").split(";", 1)[0].strip().lower()
    if media_type != "application/octet-stream" and media_type not in MEDIA_TYPES[extension]:
        raise DocumentValidationError(
            "media_type_mismatch",
            "The file extension and media type do not match.",
        )
    return extension, media_type


def validate_content_signature(path: str, extension: str) -> None:
    try:
        head = Path(path).read_bytes()[:8]
    except OSError as exc:
        raise DocumentValidationError("invalid_document", "The uploaded file could not be validated.") from exc
    if extension == ".pdf" and not head.startswith(b"%PDF-"):
        raise DocumentValidationError("invalid_pdf", "The uploaded file is not a valid PDF.")
    if extension in OFFICE_EXPECTED_ENTRY:
        if not head.startswith(b"PK"):
            raise DocumentValidationError("invalid_office_container", "The uploaded file is not a valid Office document.")
        validate_office_container(path, extension)
    if extension in {".txt", ".md", ".csv", ".json"}:
        _decode_text(Path(path).read_bytes())


def validate_office_container(path: str, extension: str) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ZIP_ENTRIES:
                raise DocumentValidationError("unsafe_office_container", "The Office document contains too many entries.")
            total_size = 0
            names: set[str] = set()
            for entry in entries:
                normalized = entry.filename.replace("\\", "/")
                parts = PurePosixPath(normalized).parts
                if normalized.startswith("/") or ".." in parts:
                    raise DocumentValidationError("unsafe_office_container", "The Office document contains an unsafe path.")
                if entry.flag_bits & 0x1:
                    raise DocumentValidationError("encrypted_document", "Encrypted documents are not supported.")
                total_size += max(0, entry.file_size)
                if total_size > MAX_ZIP_UNCOMPRESSED_BYTES:
                    raise DocumentValidationError("zip_bomb_detected", "The Office document expands beyond the safe limit.")
                compressed = max(1, entry.compress_size)
                ratio = entry.file_size / compressed
                if entry.file_size > 1_000_000 and ratio > MAX_ZIP_COMPRESSION_RATIO:
                    raise DocumentValidationError("zip_bomb_detected", "The Office document has an unsafe compression ratio.")
                lower = normalized.lower()
                if Path(lower).suffix in FORBIDDEN_ZIP_ENTRY_EXTENSIONS or "vbaproject.bin" in lower:
                    raise DocumentValidationError("unsafe_office_container", "The Office document contains executable or macro content.")
                names.add(lower)
            if OFFICE_EXPECTED_ENTRY[extension] not in names or "[content_types].xml" not in names:
                raise DocumentValidationError("invalid_office_container", "The uploaded file is not a valid Office document.")
    except DocumentValidationError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise DocumentValidationError("invalid_office_container", "The uploaded file is not a valid Office document.") from exc


def _decode_text(data: bytes) -> str:
    if b"\x00" in data:
        raise DocumentValidationError("binary_content", "The text document appears to contain binary data.")
    control_count = sum(1 for byte in data if byte < 32 and byte not in {9, 10, 13})
    if data and control_count / len(data) > 0.01:
        raise DocumentValidationError("binary_content", "The text document appears to contain binary data.")
    for encoding in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentValidationError("invalid_text_encoding", "The text document could not be decoded safely.")


def _text_sections(path: str) -> Iterable[tuple[str, str]]:
    text = _decode_text(Path(path).read_bytes())
    lines = text.splitlines()
    for start in range(0, len(lines), 80):
        block = "\n".join(lines[start:start + 80]).strip()
        if block:
            end = min(len(lines), start + 80)
            yield f"lines {start + 1}-{end}", block


def _csv_sections(path: str) -> Iterable[tuple[str, str]]:
    text = _decode_text(Path(path).read_bytes())
    reader = csv.reader(io.StringIO(text, newline=""))
    for row_number, row in enumerate(reader, start=1):
        if row_number > 2_000:
            break
        if len(row) > 100:
            raise DocumentValidationError("csv_too_wide", "The CSV contains more than 100 columns.")
        value = " | ".join(str(cell)[:1_000] for cell in row).strip()
        if value:
            yield f"row {row_number}", value


def _json_sections(path: str) -> Iterable[tuple[str, str]]:
    text = _decode_text(Path(path).read_bytes())
    try:
        payload = json.loads(text)
    except json.JSONDecodeError as exc:
        raise DocumentValidationError("invalid_json", "The JSON document is not valid JSON.") from exc
    rendered = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True)
    lines = rendered.splitlines()
    for start in range(0, len(lines), 60):
        block = "\n".join(lines[start:start + 60]).strip()
        if block:
            yield f"lines {start + 1}-{min(len(lines), start + 60)}", block


def _pdf_sections(path: str) -> tuple[list[tuple[str, str]], list[int]]:
    if not Path(path).read_bytes()[:5].startswith(b"%PDF-"):
        raise DocumentValidationError("invalid_pdf", "The uploaded file is not a valid PDF.")
    try:
        from pypdf import PdfReader

        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise DocumentValidationError("encrypted_document", "Encrypted PDFs are not supported.")
        sections: list[tuple[str, str]] = []
        page_counts: list[int] = []
        for index, page in enumerate(reader.pages[:100], start=1):
            text = str(page.extract_text() or "").strip()
            page_counts.append(len(text))
            if text:
                sections.append((f"page {index}", text))
        return sections, page_counts
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentValidationError("invalid_pdf", "The PDF could not be parsed safely.") from exc


def _docx_sections(path: str) -> Iterable[tuple[str, str]]:
    from docx import Document

    document = Document(path)
    for index, paragraph in enumerate(document.paragraphs[:10_000], start=1):
        text = paragraph.text.strip()
        if text:
            yield f"paragraph {index}", text
    for table_index, table in enumerate(document.tables[:200], start=1):
        for row_index, row in enumerate(table.rows[:2_000], start=1):
            value = " | ".join(cell.text.strip()[:2_000] for cell in row.cells).strip(" |")
            if value:
                yield f"table {table_index} row {row_index}", value
    seen_parts: set[str] = set()
    for section_index, section in enumerate(document.sections, start=1):
        for part_name, part in (("header", section.header), ("footer", section.footer)):
            for paragraph_index, paragraph in enumerate(part.paragraphs, start=1):
                text = paragraph.text.strip()
                key = f"{part_name}:{text}"
                if text and key not in seen_parts:
                    seen_parts.add(key)
                    yield f"section {section_index} {part_name} {paragraph_index}", text
    # python-docx does not expose all drawing text boxes.  The OOXML is already
    # size/path/macro validated, so bounded Word XML text is safe to parse as data.
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as archive:
        for name in sorted(archive.namelist()):
            lower = name.lower()
            if not lower.startswith("word/") or not lower.endswith(".xml"):
                continue
            root = ET.fromstring(archive.read(name))
            for index, box in enumerate(root.findall(".//w:txbxContent", namespaces), start=1):
                value = " ".join(
                    str(node.text or "").strip()
                    for node in box.findall(".//w:t", namespaces)
                    if str(node.text or "").strip()
                ).strip()
                if value:
                    yield f"text box {name} {index}", value


def _xlsx_sections(path: str) -> Iterable[tuple[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        cell_count = 0
        for worksheet in workbook.worksheets[:25]:
            for row_number, row in enumerate(
                worksheet.iter_rows(min_row=1, max_row=5_000, max_col=100, values_only=True), start=1
            ):
                cell_count += len(row)
                if cell_count > 100_000:
                    return
                values = [str(value)[:1_000] for value in row if value is not None]
                if values:
                    safe_sheet = str(worksheet.title).replace("]", "")[:80]
                    yield f"{safe_sheet} row {row_number}", " | ".join(values)
    finally:
        workbook.close()


def _pptx_sections(path: str) -> Iterable[tuple[str, str]]:
    from pptx import Presentation

    presentation = Presentation(path)
    for slide_number, slide in enumerate(presentation.slides, start=1):
        if slide_number > 200:
            break
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = str(getattr(shape, "text", "") or "").strip()
                if value:
                    parts.append(value)
            if getattr(shape, "has_table", False):
                for row_number, row in enumerate(shape.table.rows, start=1):
                    if row_number > 200:
                        break
                    value = " | ".join(cell.text.strip()[:2_000] for cell in row.cells).strip(" |")
                    if value:
                        parts.append(value)
        if parts:
            yield f"slide {slide_number}", "\n".join(parts)


EXTRACTORS: dict[str, Callable[[str], Iterable[tuple[str, str]]]] = {
    ".txt": _text_sections,
    ".md": _text_sections,
    ".csv": _csv_sections,
    ".json": _json_sections,
    ".docx": _docx_sections,
    ".xlsx": _xlsx_sections,
    ".pptx": _pptx_sections,
}


def _split_chunk(text: str) -> Iterable[str]:
    start = 0
    while start < len(text):
        end = min(len(text), start + CHUNK_CHARS)
        yield text[start:end]
        if end >= len(text):
            return
        start = max(start + 1, end - CHUNK_OVERLAP_CHARS)


def chunk_virtual_text(text: str) -> list[ExtractedChunk]:
    chunks: list[ExtractedChunk] = []
    for index, part in enumerate(_split_chunk(str(text or "")), start=1):
        if part.strip():
            chunks.append(ExtractedChunk(text=part, source=f"pasted text chunk {index}"))
    return chunks


def extract_document(path: str, extension: str) -> ExtractionResult:
    if extension in OFFICE_EXPECTED_ENTRY:
        validate_office_container(path, extension)
    limit = max_extracted_chars()
    warnings: list[str] = []
    warning_codes: list[str] = []
    page_character_counts: list[int] = []
    chunks: list[ExtractedChunk] = []
    sources: list[str] = []
    extracted_chars = 0
    try:
        if extension == ".pdf":
            sections, page_character_counts = _pdf_sections(path)
            if page_character_counts and (
                sum(page_character_counts) < max(40, len(page_character_counts) * 20)
                or sum(1 for count in page_character_counts if count < 10) / len(page_character_counts) >= 0.8
            ):
                warning_codes.append("likely_scanned_pdf")
                warnings.append(
                    "This PDF appears to be scanned or image-based. OCR was not performed; upload a text PDF or enable the isolated OCR worker."
                )
        else:
            sections = EXTRACTORS[extension](path)
        for source, raw_text in sections:
            text = re.sub(r"\r\n?", "\n", str(raw_text or "")).strip()
            if not text:
                continue
            remaining = limit - extracted_chars
            if remaining <= 0:
                warnings.append("Extraction was truncated at the configured character limit.")
                break
            if len(text) > remaining:
                text = text[:remaining]
                warnings.append("Extraction was truncated at the configured character limit.")
            if source not in sources:
                sources.append(source)
            for part in _split_chunk(text):
                if part.strip():
                    chunks.append(ExtractedChunk(text=part.strip(), source=source))
            extracted_chars += len(text)
            if extracted_chars >= limit:
                break
    except DocumentValidationError:
        raise
    except Exception as exc:
        raise DocumentValidationError("document_parse_failed", "The document could not be parsed safely.") from exc
    if not chunks and extension == ".pdf" and "likely_scanned_pdf" in warning_codes:
        return ExtractionResult(
            chunks=[], source_locators=[f"page {index}" for index in range(1, len(page_character_counts) + 1)],
            warnings=warnings, extracted_chars=0, warning_codes=warning_codes,
            page_character_counts=page_character_counts,
        )
    if not chunks:
        raise DocumentValidationError("no_extractable_text", "No extractable text was found in this document.")
    return ExtractionResult(
        chunks=chunks, source_locators=sources, warnings=warnings,
        extracted_chars=extracted_chars, warning_codes=warning_codes,
        page_character_counts=page_character_counts,
    )
